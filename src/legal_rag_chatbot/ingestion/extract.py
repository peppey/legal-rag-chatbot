"""Extract text from uploaded PDF, DOCX and TXT files."""

import io
import re
from collections.abc import Callable
from pathlib import Path

import docx
from pypdf import PdfReader
from pypdf.errors import PyPdfError

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt"}

# One unit is a page (None for formats without pages) with its paragraphs.
TextUnit = tuple[int | None, list[str]]


class ExtractionError(ValueError):
    """The file cannot be read or contains no extractable text."""


def extract_text_units(
    filename: str,
    data: bytes,
    ocr: Callable[[bytes], list[TextUnit]] | None = None,
) -> list[TextUnit]:
    """Text per page; a PDF without any text layer goes through ocr (if given)."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ExtractionError("Nur PDF, DOCX und TXT werden unterstützt.")
    reader = {".pdf": _pdf_units, ".docx": _docx_units, ".txt": _txt_units}[suffix]
    units = [(page, paragraphs) for page, paragraphs in reader(data) if paragraphs]
    if not units and suffix == ".pdf" and ocr:
        units = [(page, paragraphs) for page, paragraphs in ocr(data) if paragraphs]
    if not units:
        raise ExtractionError("Im Dokument wurde kein Text gefunden.")
    return units


def _pdf_units(data: bytes) -> list[TextUnit]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Passwortgeschützte PDFs werden nicht unterstützt.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PyPdfError as error:
        raise ExtractionError("Die PDF-Datei konnte nicht gelesen werden.") from error
    return [
        (number, [text] if text.strip() else []) for number, text in enumerate(pages, 1)
    ]


def _docx_units(data: bytes) -> list[TextUnit]:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as error:  # python-docx has no common base for corrupt input
        raise ExtractionError("Die DOCX-Datei konnte nicht gelesen werden.") from error
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [(None, [p for p in paragraphs if p.strip()])]


def _txt_units(data: bytes) -> list[TextUnit]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    return [(None, [p for p in re.split(r"\n\s*\n", text) if p.strip()])]
